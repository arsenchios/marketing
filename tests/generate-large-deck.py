# Synthetic, reproducible fixture; requires python-pptx and Pillow.
# Run from repository root. Generated binaries stay in ignored test-results/.
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import random
from pathlib import Path
out=Path('test-results/large-deck')
out.mkdir(parents=True, exist_ok=True)
photo=out/'large-photo.png'
Image.frombytes('RGB',(4608,2592),random.Random(42).randbytes(4608*2592*3)).save(photo)
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
for n in range(1,401):
 s=p.slides.add_slide(p.slide_layouts[6])
 if n==1:s.shapes.add_picture(str(photo),0,0,width=p.slide_width,height=p.slide_height)
 t=s.shapes.add_textbox(Inches(1),Inches(1),Inches(10),Inches(2));t.text=f'Slide {n} / 400'
p.save(out/'400-slides.pptx')
print((out/'400-slides.pptx').stat().st_size)
